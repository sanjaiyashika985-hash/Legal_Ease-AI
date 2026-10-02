from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape


def _filename(document_type: str, extension: str) -> str:
    name = re.sub(r"[^a-zA-Z0-9]+", "_", document_type).strip("_").lower()
    return f"{name or 'legal_document'}.{extension}"


def make_export(document: str, document_type: str, file_format: str, logo_path=None):
    file_format = file_format.lower()
    if file_format == "txt":
        return document.encode("utf-8"), "text/plain; charset=utf-8", _filename(document_type, "txt")

    if file_format == "docx":
        output = BytesIO()
        word_document = Document()
        if logo_path and Path(logo_path).is_file():
            word_document.add_picture(str(logo_path), width=Inches(1.5))
        word_document.add_heading(document_type, level=0)
        for paragraph in document.splitlines():
            word_document.add_paragraph(paragraph)
        word_document.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", _filename(document_type, "docx")

    if file_format == "pdf":
        output = BytesIO()
        pdf = SimpleDocTemplate(output, pagesize=letter, rightMargin=inch, leftMargin=inch)
        styles = getSampleStyleSheet()
        content = []
        if logo_path and Path(logo_path).is_file():
            content.extend([Image(str(logo_path), width=1.5 * inch, height=inch), Spacer(1, 12)])
        content.append(Paragraph(escape(document_type), styles["Title"]))
        for paragraph in document.splitlines():
            if paragraph.strip():
                content.append(Paragraph(escape(paragraph), styles["BodyText"]))
                content.append(Spacer(1, 6))
            else:
                content.append(Spacer(1, 6))
        pdf.build(content)
        return output.getvalue(), "application/pdf", _filename(document_type, "pdf")

    raise ValueError(f"Unsupported export format: {file_format}")